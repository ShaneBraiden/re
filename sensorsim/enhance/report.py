"""Reports written to ``temp/``: what was done, with the code and the LaTeX.

``write_method_report``
    ``temp/NAG_Method_Report.pdf``. The method, the training, the results
    and ablations, short excerpts of the actual source (pulled with
    :func:`inspect.getsource`, so they cannot drift from the code), and the
    generated LaTeX.
``write_upload_report``
    ``temp/enhance_<stem>/report.pdf``. One uploaded image: the steps, the
    before/after panels, the metrics and the LaTeX block.
``fill_rpt_template``
    A filled *copy* of ``temp/RPT_FinalReport_Template.docx``. Only the slots
    this project can honestly fill are touched: Experiment 3 (diagram output)
    gets the enhanced figure, and Experiment 10 (Overleaf) gets the generated
    LaTeX. Every other experiment's output has to come from doing that
    experiment, so those slots are left as they are.

PDFs are rendered by headless Chrome or Edge from the HTML. If neither browser
is installed, the HTML and Markdown are still written and the result says so.
"""

from __future__ import annotations

import html
import inspect
import json
import os
import shutil
import subprocess
import time
from pathlib import Path
from typing import Any

BROWSERS = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
]

TEMPLATE = Path("temp/RPT_FinalReport_Template.docx")

CSS = """
@page { size: A4; margin: 18mm 16mm; }
:root { --ink:#1b1b1a; --muted:#5f5e5a; --rule:#dddcd5; --code:#f4f3ef; --accent:#2a78d6; }
* { box-sizing: border-box; }
body { font: 10.5pt/1.5 "Segoe UI", system-ui, sans-serif; color: var(--ink); background: #fff; margin: 0; }
h1 { font-size: 20pt; margin: 0 0 4pt; letter-spacing: -0.01em; }
h2 { font-size: 13.5pt; margin: 20pt 0 6pt; padding-bottom: 3pt; border-bottom: 1px solid var(--rule); }
h3 { font-size: 11pt; margin: 14pt 0 4pt; }
p, li { margin: 0 0 6pt; }
.sub { color: var(--muted); margin-bottom: 14pt; }
.note { border-left: 3px solid var(--accent); padding: 4pt 10pt; background: #f6f9fd; margin: 8pt 0; }
figure { margin: 10pt 0; break-inside: avoid; text-align: center; }
figure img { max-width: 100%; border: 1px solid var(--rule); }
figcaption { font-size: 9pt; color: var(--muted); margin-top: 4pt; text-align: left; }
table { border-collapse: collapse; width: 100%; font-size: 9pt; margin: 6pt 0 10pt; break-inside: avoid; }
th, td { padding: 3pt 6pt; border-bottom: 1px solid var(--rule); text-align: right; font-variant-numeric: tabular-nums; }
th:first-child, td:first-child { text-align: left; }
thead th { border-bottom: 1.5px solid var(--ink); font-weight: 600; }
pre { background: var(--code); padding: 8pt 10pt; font: 8.2pt/1.38 Consolas, "Cascadia Mono", monospace;
      white-space: pre-wrap; word-break: break-word; border-radius: 4px; break-inside: auto; }
code { font-family: Consolas, monospace; font-size: 9pt; background: var(--code); padding: 0 3px; border-radius: 3px; }
.eq { font-family: "Cambria Math", Cambria, serif; font-size: 11pt; background: #fafaf8; padding: 6pt 10pt;
      border-radius: 4px; margin: 6pt 0; }
.small { font-size: 9pt; color: var(--muted); }
"""


# --------------------------------------------------------------------------
# HTML -> PDF
# --------------------------------------------------------------------------

def find_browser() -> str | None:
    env = os.environ.get("SENSORSIM_BROWSER")
    if env and Path(env).exists():
        return env
    for candidate in BROWSERS:
        if Path(candidate).exists():
            return candidate
    for name in ("chrome", "google-chrome", "chromium", "msedge"):
        found = shutil.which(name)
        if found:
            return found
    return None


def html_to_pdf(html_path: Path, pdf_path: Path, timeout: int = 120) -> Path | None:
    browser = find_browser()
    if browser is None:
        return None
    pdf_path.unlink(missing_ok=True)
    profile = html_path.parent / ".chrome-profile"
    cmd = [browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--no-first-run",
           f"--user-data-dir={profile.resolve()}", f"--print-to-pdf={pdf_path.resolve()}",
           html_path.resolve().as_uri()]
    subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=timeout, check=False)
    shutil.rmtree(profile, ignore_errors=True)
    return pdf_path if pdf_path.exists() and pdf_path.stat().st_size > 0 else None


def _page(title: str, body: str) -> str:
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<style>{CSS}</style></head><body>{body}</body></html>")


def _table_html(rows: list[dict[str, Any]], columns: list[tuple[str, str]], digits: dict[str, int] | None = None) -> str:
    digits = digits or {}
    head = "".join(f"<th>{html.escape(label)}</th>" for _, label in columns)
    body = []
    for row in rows:
        cells = []
        for key, _ in columns:
            v = row.get(key, "")
            if isinstance(v, float):
                v = "–" if v != v else f"{v:.{digits.get(key, 3)}f}"
            cells.append(f"<td>{html.escape(str(v))}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _code(obj: Any) -> str:
    return f"<pre>{html.escape(inspect.getsource(obj).rstrip())}</pre>"


def _copy(src: str | Path, dst_dir: Path) -> str:
    src = Path(src)
    dst_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst_dir / src.name)
    return f"img/{src.name}"


# --------------------------------------------------------------------------
# Per-upload report
# --------------------------------------------------------------------------

METRIC_COLUMNS = [("arm", "Output"), ("sharpness", "Sharpness"), ("noise_sigma", "Noise σ̂"),
                  ("colorfulness", "Colourfulness"), ("nr_quality", "nr_quality"),
                  ("consistency_psnr", "Consistency PSNR"), ("consistency_ssim", "Consistency SSIM")]


def write_upload_report(record: dict[str, Any], latex_block: str, out_dir: str | Path = "temp") -> dict[str, str]:
    from sensorsim.enhance import guard, noise

    folder = Path(out_dir) / f"enhance_{record['stem']}"
    img_dir = folder / "img"
    shutil.rmtree(img_dir, ignore_errors=True)
    files = {k: _copy(v, img_dir) for k, v in record["files"].items()}

    rows = [{"arm": name, **record["metrics"][key]}
            for key, name in (("input_bicubic", "Bicubic"), ("classical", "Classical"), ("nag", "NAG (ours)"))]
    steps = "".join(f"<li><b>{html.escape(s['step'])}</b>: {html.escape(s['detail'])}</li>" for s in record["steps"])
    ckpts = "".join(f"<li><code>{html.escape(k)}</code> sha256 {html.escape(str(v)[:16])}…</li>"
                    for k, v in record["checkpoints"].items())
    body = f"""
<h1>Enhancement report: {html.escape(record['name'])}</h1>
<p class="sub">SensorSim · NAG (noise-aware guarded enhancement) · {time.strftime('%d %B %Y')} ·
{record['input_shape'][1]}×{record['input_shape'][0]} px → {record['output_shape'][1]}×{record['output_shape'][0]} px
(×{record['scale']:g}) · {record['seconds']} s</p>

<h2>What was done</h2>
<p>The image was classified as a <b>{html.escape(record['kind'])}</b>. An uploaded file has no capture record,
so its noise level was estimated blindly from the image. That estimate was then carried through every stage,
as in the restoration-bank arms.</p>
<ol>{steps}</ol>
<p class="small">Checkpoints used:</p><ul class="small">{ckpts or '<li>none</li>'}</ul>

<h2>Before and after</h2>
<figure><img src="{files['comparison']}" alt="Comparison of bicubic, classical and NAG outputs">
<figcaption>Full frame of the NAG output, with the region of highest edge energy zoomed below for each
method.</figcaption></figure>
<figure><img src="{files['gate']}" alt="Wiener gate map" style="max-height:200px">
<figcaption>Wiener gate γ: the share of the generator's added detail that was kept. Black means the measured
image held nothing above the noise there, so no detail was added.</figcaption></figure>

<h2>Metrics</h2>
{_table_html(rows, METRIC_COLUMNS, {'sharpness': 1, 'noise_sigma': 2, 'colorfulness': 1, 'nr_quality': 2, 'consistency_psnr': 2, 'consistency_ssim': 4})}
<p class="small">Consistency is measured by downscaling the output back to the input size, matching its tone
to the upload per channel with an affine fit, then computing PSNR/SSIM against the upload. A high value
means the output agrees with what the file actually contained, and invented detail lowers it. The other
columns are the testbench's no-reference proxies; they are not a perceptual study.</p>

<h2>Code (short)</h2>
<h3>The Wiener gate</h3>{_code(guard.wiener_gate)}
<h3>Blind noise estimate</h3>{_code(noise.blind_sigma_map)}

<h2>LaTeX inserted into the paper</h2>
<pre>{html.escape(latex_block)}</pre>
"""
    html_path = folder / "report.html"
    html_path.write_text(_page(f"Enhancement report {record['name']}", body), encoding="utf-8")

    md = [f"# Enhancement report: {record['name']}", "",
          f"Kind: **{record['kind']}**, {record['input_shape'][1]}x{record['input_shape'][0]} -> "
          f"{record['output_shape'][1]}x{record['output_shape'][0]} px", "", "## Steps", ""]
    md += [f"{i + 1}. **{s['step']}**: {s['detail']}" for i, s in enumerate(record["steps"])]
    md += ["", "## Metrics", "", "| Output | " + " | ".join(lbl for _, lbl in METRIC_COLUMNS[1:]) + " |",
           "|---" * len(METRIC_COLUMNS) + "|"]
    for r in rows:
        md.append(f"| {r['arm']} | " + " | ".join(f"{r[k]:.3f}" for k, _ in METRIC_COLUMNS[1:]) + " |")
    md += ["", "![comparison](img/comparison.png)", "", "## LaTeX", "", "```latex", latex_block.rstrip(), "```", ""]
    (folder / "report.md").write_text("\n".join(md), encoding="utf-8")

    pdf = html_to_pdf(html_path, folder / "report.pdf")
    out = {"report_html": html_path.as_posix(), "report_md": (folder / "report.md").as_posix()}
    if pdf:
        out["report_pdf"] = pdf.as_posix()
    else:
        out["report_pdf_error"] = "no Chrome/Edge found for PDF rendering; HTML and Markdown written"
    return out


# --------------------------------------------------------------------------
# RPT course-report template
# --------------------------------------------------------------------------

def _insert_after(paragraph, new_paragraph) -> None:
    paragraph._p.addnext(new_paragraph._p)


def fill_rpt_template(
    record: dict[str, Any],
    latex_block: str,
    out_dir: str | Path = "temp",
    cover: dict[str, str | None] | None = None,
    template: str | Path = TEMPLATE,
) -> str:
    """A filled copy of the RPT template; the template itself is never written."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu, Pt

    template = Path(template)
    if not template.exists():
        raise FileNotFoundError(f"RPT template not found at {template}")
    doc = Document(str(template))
    section = doc.sections[-1]
    usable = Emu(section.page_width - section.left_margin - section.right_margin)
    cover = {k: v for k, v in (cover or {}).items() if v}

    def replace_runs(paragraph, text: str) -> None:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""

    for p in doc.paragraphs:
        if cover.get("title") and p.text.strip() == "PROJECT TITLE":
            replace_runs(p, cover["title"].upper())
        elif (cover.get("student") or cover.get("uid")) and p.text.strip().startswith("Student Name"):
            replace_runs(p, f"{cover.get('student', 'Student Name')} – {cover.get('uid', 'UID')}")
        elif (cover.get("student") or cover.get("uid")) and "</Student Name/>(UID)" in p.text:
            for run in p.runs:
                run.text = run.text.replace("</Student Name/>", cover.get("student", "Student Name")).replace(
                    "(UID)", f"({cover.get('uid', 'UID')})")

    def fill_output(match: str, intro: str, pictures: list[tuple[str, str]], code: str | None = None) -> bool:
        for p in doc.paragraphs:
            if p.text.startswith("Output:") and match in p.text:
                p.runs[-1].text = intro
                anchor = p
                for path, caption in pictures:
                    pic = doc.add_paragraph()
                    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pic.add_run().add_picture(str(path), width=usable)
                    _insert_after(anchor, pic)
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(caption)
                    r.italic = True
                    r.font.size = Pt(10)
                    _insert_after(pic, cap)
                    anchor = cap
                if code:
                    for line in code.rstrip().splitlines():
                        para = doc.add_paragraph()
                        para.paragraph_format.space_after = Pt(0)
                        run = para.add_run(line if line else " ")
                        run.font.name = "Consolas"
                        run.font.size = Pt(8)
                        _insert_after(anchor, para)
                        anchor = para
                return True
        return False

    files = record["files"]
    fill_output(
        "Draw.io",
        f"The research diagram ({record['input_shape'][1]}×{record['input_shape'][0]} px) was enhanced to "
        f"{record['output_shape'][1]}×{record['output_shape'][0]} px for print using the project's NAG pipeline "
        f"(noise-guarded Real-ESRGAN). The enhanced diagram and a zoomed comparison are shown below.",
        [(files["nag"], "Figure: enhanced research diagram (NAG output)."),
         (files["comparison"], "Figure: zoomed comparison - bicubic, classical baseline, NAG.")],
    )
    fill_output(
        "Overleaf",
        "Insert your similarity report screenshot here. The LaTeX block below was generated for the "
        "Overleaf paper; compile it on Overleaf and add the screenshot of the rendered page.",
        [],
        code=latex_block,
    )

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"RPT_FinalReport_{record['stem']}.docx"
    doc.save(str(target))
    return target.as_posix()


# --------------------------------------------------------------------------
# Method report (the whole contribution)
# --------------------------------------------------------------------------

def write_method_report(assets: dict[str, Any], out_dir: str | Path = "temp",
                        upload_dirs: list[str | Path] | None = None) -> str:
    from sensorsim.enhance import guard, latex, noise, vpdenoise
    from sensorsim.restore.learned import nag
    from sensorsim.train import nc_dce

    out_dir = Path(out_dir)
    folder = out_dir / "nag_method_report"
    img_dir = folder / "img"
    shutil.rmtree(img_dir, ignore_errors=True)
    figs = {k: _copy(v, img_dir) for k, v in assets["figures"].items()}

    main = assets["main"].reset_index().rename(columns={"pipeline": "arm"}).to_dict("records")
    abl = assets["ablation"].reset_index().rename(columns={"pipeline": "arm"}).to_dict("records")
    cols = [("arm", "Arm"), ("psnr", "PSNR"), ("ssim", "SSIM"), ("delta_e", "ΔE"), ("noise_sigma", "σ̂"),
            ("nr_quality", "nr_quality"), ("sift_inlier_ratio", "SIFT inlier")]
    dig = {"psnr": 2, "ssim": 3, "delta_e": 2, "noise_sigma": 2, "nr_quality": 2, "sift_inlier_ratio": 3}

    tests = assets["tests"]
    psnr_tests = tests[tests["metric"] == "psnr"].to_dict("records")
    test_cols = [("arm", "Arm"), ("against", "Against"), ("n", "Pairs"), ("mean_delta", "Mean Δ PSNR (dB)"),
                 ("wins", "Frames higher"), ("p_holm", "p (Holm)"), ("significant", "Significant")]

    noise_df = assets["noise"]
    noise_rows = []
    for (a, s, b, bits), g in noise_df.groupby(["alpha", "sigma_r_dn", "binning", "bit_depth"], sort=False):
        noise_rows.append({"cond": f"α={a:g}, σr={s:g} DN, b={b}, {bits}-bit",
                           "ratio": float(g["ratio"].mean()), "worst": float((g["ratio"] - 1).abs().max() * 100)})

    train_rows = [{"name": k, "status": f"diverged (epoch {v['diverged']['epoch']})" if v.get("diverged") else "trained",
                   "in_ch": v.get("in_ch"), "noise_loss": v.get("noise_loss"),
                   "epochs": v.get("config", {}).get("epochs"), "seconds": v.get("seconds"),
                   "final_total": v.get("final", {}).get("total", float("nan")),
                   "final_noise": v.get("final", {}).get("noise_metric", float("nan")),
                   "sha": str(v.get("sha256", ""))[:16]}
                  for k, v in assets["training"].items()]

    upload_html = ""
    for d in upload_dirs or []:
        rec_path = Path(d) / "metrics.json"
        if rec_path.exists():
            rec = json.loads(rec_path.read_text(encoding="utf-8"))
            _copy(rec["files"]["comparison"], img_dir / rec["id"])
            src = f"img/{rec['id']}/{Path(rec['files']['comparison']).name}"
            upload_html += (f"<figure><img src='{src}'><figcaption>Upload <code>{html.escape(rec['name'])}</code> "
                            f"({html.escape(rec['kind'])}), ×{rec['scale']:g}: bicubic, classical baseline, NAG."
                            f"</figcaption></figure>")

    generated = Path("paper/overleaf/generated")
    latex_full = latex.SECTION + "\n% ---- generated/proposed_results.tex ----\n" + \
        (generated / "proposed_results.tex").read_text(encoding="utf-8") if (generated / "proposed_results.tex").exists() else latex.SECTION

    findings = "".join(f"<li>{html.escape(_strip_tex(f))}</li>" for f in assets["findings"])
    body = f"""
<h1>NAG: Noise-Aware Guarded Enhancement</h1>
<p class="sub">What was built on top of SensorSim, how it works, the code that implements it, and the LaTeX
added to the paper · {time.strftime('%d %B %Y')}</p>

<div class="note">This report and the code it describes were prepared with AI assistance (Claude Code).
The method, the experiments and every number below can be regenerated from the repository
(<code>sensorsim train nc-dce</code>, <code>sensorsim run configs/experiments/enhance_study.yaml</code>,
<code>sensorsim paper-nag</code>). Disclose the AI assistance as your institution requires before you
submit any of this text.</div>

<h2>1. Starting point: two published methods</h2>
<p><b>Zero-DCE</b> (Guo et al., CVPR 2020) brightens a dark image with eight learned quadratic curves per
pixel. It is trained without reference images, using four losses: spatial consistency, exposure, colour
constancy and curve smoothness. None of those losses can see noise, so it lifts a noisy shadow as hard as
a clean one.</p>
<p><b>Real-ESRGAN</b> (Wang et al., ICCVW 2021) restores resolution with a generator trained on synthetic
degradations. Its strength is also its risk: where the input holds no detail, it invents some. The SensorSim
paper calls this <i>substitution</i> (hypotheses H2/H3).</p>

<h2>2. The new idea</h2>
<p>SensorSim already knows how noisy every pixel is, because its forward sensor model is Poisson–Gaussian
and every arm receives the capture record. NAG computes that noise map once, carries it through every
processing stage in closed form, and lets each stage act only where the measurement supports it.</p>
<div class="eq">var(u) = u / (α F b²) + (σᵣ / (α S b²))² + (q/α)² / 12  &nbsp;&nbsp;(+ dark current, PRNU)</div>
<p>White balance scales a channel's σ by its gain. Demosaicing scales σ by 3/4 (R, B) and √(3/4) (G). A
curve T maps σ by the symmetric secant |T(x+σ) − T(x−σ)| / 2. The resulting map was checked against the
noise measured on flat fields:</p>
{_table_html(noise_rows, [("cond", "Condition"), ("ratio", "Measured / predicted"), ("worst", "Worst error (%)")], {"ratio": 3, "worst": 1})}

<h3>NC-DCE: noise-conditioned curves (trained)</h3>
<p>This is the Zero-DCE network with the noise map as a 4th input channel, initialised from the published
weights. The noise channel's input weights start at zero, so training starts exactly at the base method.
The slope of the 8-step curve has a closed form, which gives the noise after the curve:</p>
<div class="eq">s = ∏ₙ (1 + rₙ(2·LEₙ₋₁ − 1)) &nbsp;&nbsp; σ_out = |s|·σ &nbsp;&nbsp;
L_noise = E[ w_flat · relu(|s|σ − max(σ, τ))² ] / τ²</div>
<p>The new loss stops the curves from raising noise above its input level, or above τ = 2/255 if that is
larger, in flat regions. Brightening textured regions costs no SNR and is not penalised. The exposure loss is
one-sided, so frames that are already well exposed are never darkened.</p>
<h3>VP-denoise</h3>
<p>A self-guided filter (He et al.) with a per-pixel ε = (1.5·σ_out)². It smooths where the curves amplified
noise and leaves the rest alone.</p>
<h3>G-ESRGAN: guarded super-resolution</h3>
<div class="eq">γ = clip((E_y − 20σ²) / (E_y + ε), 0, 1) &nbsp;&nbsp; x = U(y) + γ·(SR(y) − U(y)) &nbsp;&nbsp;
x ← x + U(G * (y − D(x))) × 3</div>
<p>The 4-neighbour Laplacian has noise gain 20, so γ is the Wiener fraction of local detail energy that is
signal. Generated detail is kept on real edges and removed where the measurement is only noise. Low-pass
back-projection through the recorded capture operator D (PSF, then binning) then keeps the output consistent
with what was measured.</p>

<h2>3. Training</h2>
{_table_html(train_rows, [("name", "Variant"), ("status", "Status"), ("in_ch", "Inputs"), ("noise_loss", "L_noise"), ("epochs", "Epochs"), ("seconds", "Seconds"), ("final_total", "Final loss"), ("final_noise", "Final L_noise (measured)"), ("sha", "sha256")], {"final_total": 3, "final_noise": 3, "seconds": 0})}
{f'<figure><img src="{figs["nag_training.png"]}"><figcaption>Training curves. The right panel is the noise-amplification term, measured for every variant even when it is not optimised.</figcaption></figure>' if "nag_training.png" in figs else ""}

<h2>4. Results</h2>
<p>The study uses the same corpus, sweep, seed and metrics as the paper's pilot: {assets['observations']:,} observations,
21 frames per arm. The pilot's own table was not touched.</p>
<h3>Main comparison (means over the sweep)</h3>
{_table_html(main, cols, dig)}
<h3>Ablations</h3>
{_table_html(abl, cols, dig)}
<h3>Planned paired tests (PSNR, Wilcoxon, Holm-corrected)</h3>
{_table_html(psnr_tests, test_cols, {"mean_delta": 2, "p_holm": 4})}
<h3>What the numbers say</h3>
<ul>{findings}</ul>
<figure><img src="{figs['nag_curves.png']}"><figcaption>PSNR and SIFT inlier ratio across the sweep.</figcaption></figure>
<figure><img src="{figs['nag_crops.png']}"><figcaption>Matched crops at the most severe condition.</figcaption></figure>
{('<h3>Uploaded images</h3>' + upload_html) if upload_html else ''}

<h2>5. Code, in short</h2>
<p class="small">These excerpts are pulled from the source when the report is built, so they match the code
that produced the numbers.</p>
<h3>Noise variance from the capture record — <code>sensorsim/enhance/noise.py</code></h3>{_code(noise.variance_linear)}
<h3>Closed-form curve slope — <code>sensorsim/enhance/noise.py</code></h3>{_code(noise.curve_slope_torch)}
<h3>The new training loss — <code>sensorsim/train/nc_dce.py</code></h3>{_code(nc_dce.loss_noise)}
<h3>VP-denoise — <code>sensorsim/enhance/vpdenoise.py</code></h3>{_code(vpdenoise.guided_filter_varying)}
<h3>Wiener gate and back-projection — <code>sensorsim/enhance/guard.py</code></h3>{_code(guard.wiener_gate)}{_code(guard.back_project)}
<h3>The NAG arm — <code>sensorsim/restore/learned/nag.py</code></h3>{_code(nag.NoiseAwareGuarded.restore)}

<h2>6. LaTeX added to the paper</h2>
<p class="small">These blocks sit in <code>paper/overleaf/main.tex</code> between <code>%% AUTOGEN:nag-*</code>
markers; a rerun replaces them in place. Every number is a macro from <code>generated/proposed_numbers.tex</code>.</p>
<pre>{html.escape(latex_full)}</pre>
<h3>generated/proposed_numbers.tex</h3>
<pre>{html.escape((generated / 'proposed_numbers.tex').read_text(encoding='utf-8') if (generated / 'proposed_numbers.tex').exists() else '')}</pre>

<h2>7. How to reproduce</h2>
<pre>sensorsim weights fetch                       # pinned Real-ESRGAN / Zero-DCE checkpoints
sensorsim train nc-dce                        # NC-DCE + two ablations
sensorsim run configs/experiments/enhance_study.yaml
sensorsim paper-nag                           # tables, figures, macros, paper blocks, this report
sensorsim enhance image.png                   # one upload: enhance, figure into the paper, reports
sensorsim dashboard --upload                  # the same, from the browser</pre>
"""
    html_path = folder / "NAG_Method_Report.html"
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(_page("NAG Method Report", body), encoding="utf-8")
    pdf = html_to_pdf(html_path, out_dir / "NAG_Method_Report.pdf")
    return (pdf or html_path).as_posix()


def _strip_tex(text: str) -> str:
    """Generated findings are LaTeX; make them readable as plain text."""
    import re

    replacements = {r"\GESR": "G-ESRGAN", r"\NCDCE": "NC-DCE", r"\NAGarm": "NAG", r"\RESR": "Real-ESRGAN",
                    r"\ZDCE": "Zero-DCE", r"\ISP": "control ISP", r"\hat\sigma": "σ̂", r"\mathcal{L}_{\mathrm{noise}}": "L_noise",
                    r"\sigma": "σ", "~": " ", r"\ ": " "}
    for k, v in replacements.items():
        text = text.replace(k, v)
    text = re.sub(r"\\ensuremath\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\texttt\{([^}]*)\}", r"\1", text)
    return text.replace("$", "").replace(r"\_", "_").replace(r"\%", "%")
